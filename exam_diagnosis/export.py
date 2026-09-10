"""Report export to PDF and Word.

Both formats are built from one intermediate block list (`report_blocks`) rather than from the
markdown/HTML renderers in reporting.py, so PDF and DOCX can't drift apart on which sections a
report contains, and each format only has to know how to draw a heading, a paragraph and a list.

PDF goes through PyMuPDF's Story engine — already a dependency for reading uploads — instead of a
second PDF toolkit. Its built-in Droid Sans Fallback covers Chinese out of the box, which is the
one thing most pure-Python PDF libraries make painful.
"""
from __future__ import annotations

import io
from html import escape
from typing import Any

from .reporting import PRIORITY_LABELS, SEVERITY_LABELS, answer_or_withheld, text

REPORT_TITLE = 'Exam Diagnosis 报告'

# One block is a tuple: ('heading', level, text) | ('paragraph', text) | ('bullets', [text], level) |
# ('numbered', [text]). `level` on bullets nests evidence/examples under their cluster.
Block = tuple


def _priority(value: Any) -> str:
    return PRIORITY_LABELS.get(value, '未评估')


def report_blocks(report: dict[str, Any]) -> list[Block]:
    blocks: list[Block] = [('heading', 1, REPORT_TITLE), ('paragraph', f"错题数：{report.get('wrong_question_count', 0)}")]
    versions = report.get('exam_versions')
    if versions:
        rendered = '；'.join('{} v{}'.format(text(item.get('name')), item.get('version')) for item in versions)
        blocks.append(('paragraph', f'分析依据的考试解析版本：{rendered}'))
    summary = report.get('summary')
    if summary:
        blocks.extend([('heading', 2, '概述'), ('paragraph', summary)])
    flags = report.get('quality_flags')
    if flags:
        blocks.append(('heading', 2, '结论可信度提示'))
        blocks.append(('bullets', [f"[{SEVERITY_LABELS.get(flag.get('severity'), text(flag.get('severity')))}] {text(flag.get('message'))}" for flag in flags], 1))
    scope = report.get('scope_coverage')
    if scope:
        blocks.append(('heading', 2, '课程范围与资料覆盖范围'))
        blocks.append(('bullets', [
            f"课程范围知识点：{scope.get('curriculum_point_count', 0)}",
            f"本次资料考查到：{scope.get('assessed_point_count', 0)}（覆盖率 {scope.get('coverage_ratio', 0)}）",
            f"本次资料未考查：{scope.get('not_assessed_point_count', 0)}",
            f"未考查章节：{'、'.join(scope.get('not_assessed_chapters', [])) or '无'}",
            text(scope.get('note')),
        ], 1))
    blocks.append(('heading', 2, '错因专题'))
    for cluster in report.get('clusters', []):
        blocks.append(('bullets', [f"{cluster.get('root_cause', '待人工确认')}（优先级 {cluster.get('priority', 0)}）：来源题目 {', '.join(cluster.get('question_ids', []))}；规避建议：{text(cluster.get('avoidance_advice'))}"], 1))
        nested = [f"原题证据：{evidence.get('original_source', {})}" for evidence in cluster.get('evidence', [])]
        nested += [f"题干：{text(example.get('question_text'))}；学生答案：{text(example.get('student_answer'))}" for example in cluster.get('question_examples', [])]
        if nested:
            blocks.append(('bullets', nested, 2))
    blocks.append(('heading', 2, '知识树'))
    for chapter, points in report.get('knowledge_tree', {}).items():
        blocks.append(('heading', 3, chapter))
        blocks.append(('bullets', [f"{name}：错题 {stats.get('wrong_count', 0)}/{stats.get('question_count', 0)}；错误率 {stats.get('error_rate', 0)}；复习优先级 {_priority(stats.get('review_priority'))}" for name, stats in points.items()], 1))
    review_order = report.get('review_order')
    if review_order:
        blocks.append(('heading', 2, '复习优先顺序'))
        blocks.append(('numbered', [f"{text(item.get('chapter'))} / {text(item.get('name'))}：优先级 {_priority(item.get('review_priority'))}（错题 {item.get('wrong_count', 0)}/{item.get('question_count', 0)}）" for item in review_order]))
    blocks.append(('heading', 2, '错题集'))
    blocks.append(('bullets', [f"题干：{text(item.get('question_text'))}；学生答案：{text(item.get('student_answer'))}；错因：{text(item.get('root_cause'))}；规避建议：{text(item.get('avoidance_advice'))}" for item in report.get('wrong_question_collection', [])], 1))
    blocks.append(('heading', 2, '验证题'))
    blocks.append(('bullets', [f"题目：{text(question.get('question_text', question.get('prompt')))}；答案要点：{answer_or_withheld(question)}" for question in report.get('verification_questions', [])], 1))
    return blocks


# --- PDF -------------------------------------------------------------------------------------

PDF_CSS = """
body { font-family: sans-serif; font-size: 10.5pt; line-height: 1.5; color: #1f1f1f; }
h1 { font-size: 20pt; margin: 0 0 6pt 0; border-bottom: 1.5pt solid #c7391f; padding-bottom: 4pt; }
h2 { font-size: 14pt; color: #c7391f; margin: 14pt 0 4pt 0; }
h3 { font-size: 11.5pt; margin: 8pt 0 2pt 0; }
p { margin: 0 0 4pt 0; }
ul, ol { margin: 0 0 4pt 0; padding-left: 16pt; }
ul.nested { padding-left: 32pt; color: #555; }
li { margin: 0 0 2pt 0; }
"""


def blocks_to_html(blocks: list[Block]) -> str:
    parts: list[str] = []
    for block in blocks:
        kind = block[0]
        if kind == 'heading':
            parts.append(f'<h{block[1]}>{escape(block[2])}</h{block[1]}>')
        elif kind == 'paragraph':
            parts.append(f'<p>{escape(block[1])}</p>')
        elif kind == 'bullets':
            css_class = ' class="nested"' if block[2] > 1 else ''
            parts.append(f'<ul{css_class}>' + ''.join(f'<li>{escape(item)}</li>' for item in block[1]) + '</ul>')
        elif kind == 'numbered':
            parts.append('<ol>' + ''.join(f'<li>{escape(item)}</li>' for item in block[1]) + '</ol>')
    return '<body>' + ''.join(parts) + '</body>'


def render_pdf(report: dict[str, Any]) -> bytes:
    return render_blocks_pdf(report_blocks(report))


def render_blocks_pdf(blocks: list[Block]) -> bytes:
    import pymupdf

    html = blocks_to_html(blocks)
    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    story = pymupdf.Story(html=html, user_css=PDF_CSS)
    page_rect = pymupdf.paper_rect('a4')
    content_rect = page_rect + (48, 48, -48, -56)
    more = True
    while more:
        device = writer.begin_page(page_rect)
        more, _ = story.place(content_rect)
        story.draw(device)
        writer.end_page()
    writer.close()
    return _subset_fonts(buffer.getvalue())


def _subset_fonts(pdf_bytes: bytes) -> bytes:
    """Story embeds the whole Droid Sans Fallback face (~3.5 MB) for even one Chinese character.
    Subsetting needs fontTools; if it's missing the PDF is still correct, just fat."""
    try:
        import fontTools  # noqa: F401
        import pymupdf
    except ImportError:
        return pdf_bytes
    document = pymupdf.open(stream=pdf_bytes, filetype='pdf')
    try:
        document.subset_fonts()
        return document.tobytes(garbage=4, deflate=True)
    except Exception:
        return pdf_bytes
    finally:
        document.close()


# --- Word ------------------------------------------------------------------------------------

DOCX_EAST_ASIAN_FONT = '宋体'


def render_docx(report: dict[str, Any]) -> bytes:
    return render_blocks_docx(report_blocks(report))


def render_blocks_docx(blocks: list[Block]) -> bytes:
    from docx import Document
    from docx.oxml.ns import qn

    document = Document()
    # python-docx's default template is Calibri with no East Asian font set; Word would pick one
    # anyway, but naming it keeps the file rendering the same on every machine that opens it.
    for style_name in ('Normal', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Bullet 2', 'List Number'):
        rpr = document.styles[style_name].element.get_or_add_rPr()
        rpr.get_or_add_rFonts().set(qn('w:eastAsia'), DOCX_EAST_ASIAN_FONT)
    for block in blocks:
        kind = block[0]
        if kind == 'heading':
            document.add_heading(block[2], level=block[1])
        elif kind == 'paragraph':
            document.add_paragraph(block[1])
        elif kind == 'bullets':
            for item in block[1]:
                document.add_paragraph(item, style='List Bullet 2' if block[2] > 1 else 'List Bullet')
        elif kind == 'numbered':
            for item in block[1]:
                document.add_paragraph(item, style='List Number')
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


EXPORT_MEDIA_TYPES = {
    'pdf': 'application/pdf',
    'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
}
