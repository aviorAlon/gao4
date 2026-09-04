"""Report export to PDF / Word: one block model feeding both renderers, plus the download route."""
import io

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient

import exam_diagnosis.app as app_module
from exam_diagnosis.export import EXPORT_MEDIA_TYPES, blocks_to_html, render_docx, render_pdf, report_blocks
from exam_diagnosis.reporting import WITHHELD_ANSWER


@pytest.fixture
def report():
    return {
        'wrong_question_count': 2,
        'exam_versions': [{'name': '期中', 'version': 3}],
        'summary': '这次两道错题都集中在函数定义域。',
        'quality_flags': [{'severity': 'warning', 'message': '有 1 页未能自动解析'}],
        'scope_coverage': {'curriculum_point_count': 40, 'assessed_point_count': 4, 'coverage_ratio': 0.1, 'not_assessed_point_count': 36, 'not_assessed_chapters': ['导数'], 'note': '仅覆盖本次资料'},
        'clusters': [{'root_cause': '条件遗漏', 'priority': 1, 'question_ids': ['q1', 'q2'], 'avoidance_advice': '先写定义域', 'evidence': [{'original_source': {'asset_id': 'a1'}}], 'question_examples': [{'question_text': '求 f(x) 定义域', 'student_answer': 'x>0'}]}],
        'knowledge_tree': {'函数': {'定义域': {'wrong_count': 2, 'question_count': 2, 'error_rate': 1.0, 'review_priority': 'high'}}},
        'review_order': [{'chapter': '函数', 'name': '定义域', 'review_priority': 'high', 'wrong_count': 2, 'question_count': 2}],
        'wrong_question_collection': [{'question_text': '求 f(x) 定义域', 'student_answer': 'x>0', 'root_cause': '条件遗漏', 'avoidance_advice': '先写定义域'}],
        'verification_questions': [
            {'question_text': '变式：求 g(x) 定义域', 'answer': 'x≠2'},
            {'question_text': '未复核的题', 'answer': '秘密', 'answer_withheld': True, 'withheld_reason': ['答案未复核']},
        ],
    }


# --- block model -----------------------------------------------------------------------------


def test_report_blocks_cover_every_section_in_order(report):
    blocks = report_blocks(report)
    headings = [b[2] for b in blocks if b[0] == 'heading']
    assert headings == ['Exam Diagnosis 报告', '概述', '结论可信度提示', '课程范围与资料覆盖范围', '错因专题', '知识树', '函数', '复习优先顺序', '错题集', '验证题']
    assert ('paragraph', '错题数：2') in blocks
    assert ('paragraph', '分析依据的考试解析版本：期中 v3') in blocks
    nested = [b for b in blocks if b[0] == 'bullets' and b[2] == 2]
    assert nested and '原题证据' in nested[0][1][0] and '学生答案：x>0' in nested[0][1][1]
    numbered = next(b for b in blocks if b[0] == 'numbered')
    assert numbered[1] == ['函数 / 定义域：优先级 高（错题 2/2）']


def test_report_blocks_skip_optional_sections_when_absent():
    headings = [b[2] for b in report_blocks({'wrong_question_count': 0}) if b[0] == 'heading']
    assert headings == ['Exam Diagnosis 报告', '错因专题', '知识树', '错题集', '验证题']


def test_withheld_verification_answers_stay_withheld_in_exports(report):
    blocks = report_blocks(report)
    verification = blocks[-1][1]
    assert verification[0].endswith('答案要点：x≠2')
    assert WITHHELD_ANSWER in verification[1] and '秘密' not in verification[1]
    assert '秘密' not in blocks_to_html(blocks)
    pdf_text = ''.join(page.get_text() for page in pymupdf.open(stream=render_pdf(report), filetype='pdf'))
    assert '秘密' not in pdf_text
    docx_text = '\n'.join(p.text for p in Document(io.BytesIO(render_docx(report))).paragraphs)
    assert '秘密' not in docx_text


def test_blocks_to_html_escapes_user_content():
    html = blocks_to_html([('paragraph', '<script>alert(1)</script>'), ('bullets', ['a & b'], 1)])
    assert '<script>' not in html and '&lt;script&gt;' in html and 'a &amp; b' in html


# --- PDF -------------------------------------------------------------------------------------


def test_render_pdf_produces_readable_chinese_pdf(report):
    data = render_pdf(report)
    assert data.startswith(b'%PDF')
    document = pymupdf.open(stream=data, filetype='pdf')
    text = ''.join(page.get_text() for page in document)
    for expected in ('Exam Diagnosis 报告', '错题数：2', '这次两道错题都集中在函数定义域', '条件遗漏（优先级 1）', '函数 / 定义域：优先级 高', '变式：求 g(x) 定义域'):
        assert expected in text
    # Chinese must actually be drawn, not silently dropped as missing glyphs: a page with this
    # much text is nowhere near blank.
    pixmap = document[0].get_pixmap(dpi=40)
    dark = sum(1 for i in range(0, len(pixmap.samples), pixmap.n) if pixmap.samples[i] < 128)
    assert dark > 500


def test_render_pdf_subsets_the_cjk_font(report):
    pytest.importorskip('fontTools')
    # Without subsetting Story embeds the entire Droid Sans Fallback face (~3.5 MB) per document.
    assert len(render_pdf(report)) < 600_000


def test_render_pdf_paginates_long_reports(report):
    report['wrong_question_collection'] = [{'question_text': f'第 {i} 题的很长的题干' * 6, 'student_answer': 'x', 'root_cause': '计算错误', 'avoidance_advice': '检查'} for i in range(120)]
    document = pymupdf.open(stream=render_pdf(report), filetype='pdf')
    assert len(document) > 1
    assert '第 119 题' in document[-1].get_text() or '第 119 题' in ''.join(p.get_text() for p in document)


# --- Word ------------------------------------------------------------------------------------


def test_render_docx_uses_heading_and_list_styles(report):
    document = Document(io.BytesIO(render_docx(report)))
    paragraphs = document.paragraphs
    assert paragraphs[0].text == 'Exam Diagnosis 报告' and paragraphs[0].style.name == 'Heading 1'
    by_style = {}
    for paragraph in paragraphs:
        by_style.setdefault(paragraph.style.name, []).append(paragraph.text)
    assert '概述' in by_style['Heading 2'] and '函数' in by_style['Heading 3']
    assert any(t.startswith('条件遗漏（优先级 1）') for t in by_style['List Bullet'])
    assert any(t.startswith('原题证据') for t in by_style['List Bullet 2'])
    assert by_style['List Number'] == ['函数 / 定义域：优先级 高（错题 2/2）']
    assert '这次两道错题都集中在函数定义域。' in by_style['Normal']


def test_render_docx_names_an_east_asian_font():
    from docx.oxml.ns import qn
    document = Document(io.BytesIO(render_docx({'wrong_question_count': 0})))
    for style_name in ('Normal', 'Heading 2', 'List Bullet'):
        rfonts = document.styles[style_name].element.rPr.rFonts
        assert rfonts.get(qn('w:eastAsia')) == '宋体'


# --- route -----------------------------------------------------------------------------------


def _run_with_report(client):
    submission = client.post('/api/submissions', json={'stage': 'high', 'subject': '数学', 'province': '广东', 'school_grade': '高二', 'academic_year_label': '25-26'}).json()
    exam = client.post(f"/api/submissions/{submission['id']}/exams", json={'name': '月考'}).json()
    client.post(f"/api/exams/{exam['id']}/confirm", json={'questions': [{'question_number': '1', 'question_text': '函数定义域', 'is_wrong_question': True, 'original_question_available': True}]})
    run = client.post(f"/api/submissions/{submission['id']}/analysis-runs", json={'exam_ids': [exam['id']]}).json()
    client.post(f"/api/analysis-runs/{run['id']}/confirm-start")
    client.get(f"/api/analysis-runs/{run['id']}/report")
    return run['id']


def test_export_route_serves_pdf_and_docx_as_downloads():
    client = TestClient(app_module.app)
    rid = _run_with_report(client)
    pdf = client.get(f'/api/reports/{rid}/export?format=pdf')
    assert pdf.status_code == 200
    assert pdf.headers['content-type'] == EXPORT_MEDIA_TYPES['pdf']
    assert pdf.headers['content-disposition'] == f'attachment; filename="exam-diagnosis-report-{rid[:8]}.pdf"'
    assert pdf.content.startswith(b'%PDF')
    assert '函数定义域' in ''.join(p.get_text() for p in pymupdf.open(stream=pdf.content, filetype='pdf'))
    docx = client.get(f'/api/reports/{rid}/export?format=docx')
    assert docx.status_code == 200
    assert docx.headers['content-type'] == EXPORT_MEDIA_TYPES['docx']
    assert docx.headers['content-disposition'].endswith(f'{rid[:8]}.docx"')
    assert any('函数定义域' in p.text for p in Document(io.BytesIO(docx.content)).paragraphs)


def test_export_route_still_rejects_unknown_formats_and_missing_reports():
    client = TestClient(app_module.app)
    rid = _run_with_report(client)
    assert client.get(f'/api/reports/{rid}/export?format=xlsx').status_code == 422
    assert client.get('/api/reports/nope/export?format=pdf').status_code == 404
